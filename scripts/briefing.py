"""매일 아침 브리핑 PPT(3장)를 만든다: 보유 주식 변동 → 변동 사유 → 미국 뉴스.

- 시세: data/history.json 의 최신 기록 (update.py 가 먼저 실행돼야 함)
- 변동 사유: 많이 움직인 종목마다 Google 뉴스(한국어) 헤드라인을 붙인다. AI 추론이 아니라 헤드라인 모음이다.
- 미국 뉴스: 연합뉴스 시장 RSS + 한국경제 국제 RSS 에서 미국 관련 기사 제목과 요약문을 가져온다.
- 결과: briefing/briefing.pptx, briefing/meta.json, briefing/index.html (슬라이드 이미지는 워크플로에서 PDF→PNG 로 만든다)
"""
import html
import json
import re
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
HISTORY = ROOT / "data" / "history.json"
OUT = ROOT / "briefing"
KST = timezone(timedelta(hours=9))
UA = {"User-Agent": "Mozilla/5.0 (portfolio-map briefing)"}
DAYS = "월화수목금토일"

FONT = "Malgun Gothic"
INK, MUTED, FAINT = RGBColor(0x16, 0x18, 0x1D), RGBColor(0x6B, 0x6E, 0x76), RGBColor(0x9A, 0x9C, 0xA3)
CARD, LINE, WHITE = RGBColor(0xF4, 0xF5, 0xF7), RGBColor(0xE3, 0xE4, 0xE8), RGBColor(0xFF, 0xFF, 0xFF)
UP_T, UP_B = RGBColor(0xC8, 0x32, 0x2F), RGBColor(0xE2, 0x4B, 0x4A)
DN_T, DN_B = RGBColor(0x1F, 0x5F, 0xC4), RGBColor(0x37, 0x8A, 0xDD)

# 변동 사유 검색어 (holdings.json 의 id → Google 뉴스 검색어). 없으면 종목명으로 검색한다.
NEWS_QUERY = {
    "kodex-nasdaq100": "나스닥 지수", "irp-kodex-nasdaq100": "나스닥 지수", "qqqm": "나스닥 지수",
    "tqqq": "나스닥 지수", "qld": "나스닥 지수", "irp-ace-nasdaq-bond": "나스닥 지수",
    "rise-sp500": "S&P500 지수", "irp-kodex-sp500": "S&P500 지수", "schd": "SCHD 배당",
    "soxl": "필라델피아 반도체지수", "dram": "마이크론 주가", "ram": "마이크론 주가",
    "googl": "알파벳 구글 주가", "hood": "로빈후드 주가", "orcl": "오라클 주가", "pltr": "팔란티어 주가",
    "samsung": "삼성전자 주가", "lgdisplay": "LG디스플레이 주가", "irp-rise-semi-bond": "삼성전자 SK하이닉스 주가",
}
US_WORDS = r"뉴욕|미국|美|연준|Fed|FOMC|나스닥|S&P|다우|월가|파월|트럼프|국채|달러|엔비디아|애플|테슬라|빅테크|마이크론|오라클|인플레|고용|관세"
SKIP_WORDS = r"축구|야구|농구|골프|올림픽|월드컵|배우|가수|드라마|영화|날씨|사망|살해|화재|포토|\[사진\]|\[영상\]"
HOLD_WORDS = [  # 미국 뉴스에 붙이는 '내 종목 관련' 표시
    (r"나스닥|빅테크|기술주", "나스닥100"), (r"S&P|다우|뉴욕증시|월가", "S&P500"), (r"반도체|엔비디아|마이크론|메모리|HBM|D램", "반도체"),
    (r"삼성전자", "삼성전자"), (r"오라클", "오라클"), (r"팔란티어", "팔란티어"), (r"구글|알파벳", "알파벳"), (r"로빈후드", "로빈후드"),
]


# ---------- data ----------
def fetch(url, timeout=20):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def rss(url):
    for attempt in range(3):
        try:
            return ET.fromstring(fetch(url)).iter("item")
        except Exception as e:
            if attempt == 2:
                print(f"  ! RSS 실패 {url}: {e}")
                return []
            time.sleep(3)


def clean(t):
    t = html.unescape(re.sub(r"<[^>]+>", "", t or "")).replace("\xa0", " ")
    return re.sub(r"\s+", " ", t).strip()


def pub_time(it):
    try:
        return parsedate_to_datetime(it.findtext("pubDate")).astimezone(KST)
    except Exception:
        return None


def portfolio():
    snaps = sorted(json.loads(HISTORY.read_text(encoding="utf-8"))["snapshots"], key=lambda s: s["date"])
    snap = snaps[-1]
    fx = snap["fx"]["usdkrw"]
    merged = {}  # 계좌만 다른 같은 종목은 합친다
    for h in snap["holdings"]:
        v = h["qty"] * h["price"] * (fx if h["currency"] == "USD" else 1)
        if v <= 0:
            continue
        prev = v / (1 + (h.get("chgPct") or 0) / 100)
        m = merged.setdefault(h["ticker"], {**h, "value": 0.0, "prev": 0.0, "ids": []})
        m["value"] += v
        m["prev"] += prev
        m["ids"].append(h["id"])
    rows = list(merged.values())
    total, prev_total = sum(r["value"] for r in rows), sum(r["prev"] for r in rows)
    for r in rows:
        r["diff"] = r["value"] - r["prev"]
        r["weight"] = r["value"] / total * 100
    rows.sort(key=lambda r: r["chgPct"], reverse=True)
    return snap, rows, total, prev_total


def strip_source(title, src):
    """Google 뉴스 제목 끝의 ' - 언론사' (가끔 두 번 붙음) 와 남는 '|' 같은 꼬리를 뗀다."""
    tail = r"[\s|\-–·:]+$"
    title = re.sub(tail, "", title)
    while src and title.endswith(src) and len(title) > len(src) + 5:
        title = re.sub(tail, "", title[: -len(src)])
    if not src:
        title = re.sub(r"\s+-\s+[^-]{1,25}$", "", title)
    return title.strip()


def mover_news(row, used):
    """종목 하나의 최근 헤드라인 2개. used 는 다른 카드에서 이미 쓴 기사(중복 방지)."""
    q = next((NEWS_QUERY[i] for i in row["ids"] if i in NEWS_QUERY), row["name"])
    url = "https://news.google.com/rss/search?q=" + urllib.parse.quote(f"{q} when:2d") + "&hl=ko&gl=KR&ceid=KR:ko"
    out = []
    for it in rss(url):
        src = clean(it.findtext("source"))
        title = strip_source(clean(it.findtext("title")), src)
        key = re.sub(r"\W", "", title)[:18]
        if not title or key in used or re.search(SKIP_WORDS, title) or re.search(r"\|.*\||Chg%|가격\s*:", title):
            continue
        used.add(key)
        out.append({"title": title, "source": src, "link": clean(it.findtext("link"))})
        if len(out) == 2:
            break
    return q, out


def article_summary(x):
    """기사 페이지에서 요약을 뽑는다: 연합뉴스는 '세 줄 요약', 한국경제는 본문 첫 문장."""
    try:
        page = fetch(x["link"], timeout=12).decode("utf-8", "ignore")
    except Exception:
        return []
    m = re.search(r'<article class="story-summary">(.*?)</article>', page, re.S)
    if m:  # 연합뉴스
        lines = [clean(p) for p in re.findall(r"<p[^>]*>(.*?)</p>", m.group(1), re.S)]
        return [l for l in lines if l][:3]
    m = re.search(r'id="articletxt"[^>]*>(.*?)</div>', page, re.S)
    if m:  # 한국경제: 사진 설명·기자 서명을 뺀 첫 문단
        for part in re.split(r"<br\s*/?>|</p>", m.group(1)):
            t = clean(part)
            if len(t) >= 25 and not re.match(r"^(사진|그래픽|\[)", t) and "@" not in t:
                return [t]
    return []


def us_news(now, limit=5):
    items = []
    for it in rss("https://www.yna.co.kr/rss/market.xml"):
        items.append({"title": clean(it.findtext("title")), "source": "연합뉴스", "link": clean(it.findtext("link")), "time": pub_time(it)})
    for it in rss("https://www.hankyung.com/feed/international"):
        items.append({"title": clean(it.findtext("title")), "source": "한국경제", "link": clean(it.findtext("link")), "time": pub_time(it)})
    fresh = [x for x in items if x["title"] and x["time"] and now - x["time"] < timedelta(hours=26)
             and re.search(US_WORDS, x["title"]) and not re.search(SKIP_WORDS, x["title"])]
    # 뉴욕증시 시황 기사를 먼저, 그다음 최신순
    fresh.sort(key=lambda x: x["time"], reverse=True)
    fresh.sort(key=lambda x: 0 if re.search(r"뉴욕증시|뉴욕 증시|월가", x["title"]) else 1)
    picked, seen = [], set()
    for x in fresh:
        key = re.sub(r"\W", "", x["title"])[:14]
        if key in seen:
            continue
        seen.add(key)
        x["summary"] = article_summary(x)
        if not x["summary"]:  # 요약을 못 뽑은 기사는 건너뛴다
            continue
        picked.append(x)
        if len(picked) == limit:
            break
    for x in picked:
        x["tags"] = [tag for pat, tag in HOLD_WORDS if re.search(pat, x["title"] + " " + " ".join(x["summary"]))][:2]
    return picked


# ---------- pptx helpers ----------
def cut(t, n):
    t = t or ""
    return t if len(t) <= n else t[: n - 1].rstrip() + "…"


def man(v, sign=False):
    s = f"{abs(v) / 10000:,.0f}만원" if abs(v) >= 10000 or v == 0 else f"{abs(v):,.0f}원"
    return (("+" if v > 0 else "−" if v < 0 else "") if sign else "") + s


def pct(v):
    return f"{'+' if v > 0 else '−' if v < 0 else ''}{abs(v):.2f}%"


def tone(v):
    return (UP_T, UP_B) if v > 0 else (DN_T, DN_B) if v < 0 else (MUTED, LINE)


def box(slide, x, y, w, h, fill=None, line=None, radius=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    style = shp._element.find(qn("p:style"))  # 기본 테마의 그림자·윤곽선 제거
    if style is not None:
        shp._element.remove(style)
    return shp


def text(slide, x, y, w, h, runs, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: 문자열 또는 [(문자열, {size,color,bold}) ...] 의 줄 목록"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [[(runs, {})]]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for s, st in (line if isinstance(line, list) else [(line, {})]):
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.color.rgb = st.get("color", color)
            rpr = r._r.get_or_add_rPr()  # 한글(동아시아) 글꼴도 같은 글꼴로: 쉼표·괄호 앞 공백 방지
            rpr.find(qn("a:latin")).addnext(rpr.makeelement(qn("a:ea"), {"typeface": FONT}))
    return tb


def header(slide, title, page, when):
    text(slide, 0.6, 0.42, 8, 0.6, title, size=28, bold=True)
    text(slide, 7.4, 0.5, 5.33, 0.4, [[(when, {}), (f"   {page} / 3", {"color": FAINT})]], size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    box(slide, 0.6, 1.08, 12.13, 0.02, fill=LINE)


# ---------- slides ----------
def slide_changes(prs, when, rows, total, prev_total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "보유 주식 변동", 1, when)
    d = total - prev_total
    p = d / prev_total * 100 if prev_total else 0
    ups, dns = sum(r["chgPct"] > 0 for r in rows), sum(r["chgPct"] < 0 for r in rows)
    kpis = [("총 평가금액", man(total), INK), ("전일 대비", man(d, True), tone(d)[0]),
            ("등락률", pct(p), tone(p)[0]), ("상승 · 하락 종목", f"{ups} / {dns}", INK)]
    for i, (lbl, val, col) in enumerate(kpis):
        x = 0.6 + i * 3.08
        box(s, x, 1.32, 2.9, 1.0, fill=CARD, radius=0.12)
        text(s, x + 0.22, 1.44, 2.5, 0.3, lbl, size=11, color=MUTED)
        text(s, x + 0.22, 1.72, 2.5, 0.5, val, size=22, bold=True, color=col)

    top, bottom = 2.62, 7.1
    n = len(rows)
    rh = min(0.34, (bottom - top) / max(n, 1))
    fs = 12 if rh >= 0.3 else 10.5
    mx = max((abs(r["chgPct"]) for r in rows), default=1) or 1
    zero, half = 6.55, 2.9  # 막대 영역: 3.65 ~ 9.45
    text(s, 0.6, top - 0.28, 2.8, 0.25, "종목", size=10, color=FAINT)
    text(s, 9.65, top - 0.28, 1.2, 0.25, "등락률", size=10, color=FAINT, align=PP_ALIGN.RIGHT)
    text(s, 10.95, top - 0.28, 1.78, 0.25, "평가금액 변동", size=10, color=FAINT, align=PP_ALIGN.RIGHT)
    box(s, zero, top, 0.012, rh * n, fill=LINE)
    for i, r in enumerate(rows):
        y = top + i * rh
        t_col, b_col = tone(r["chgPct"])
        text(s, 0.6, y, 2.95, rh, cut(r.get("label") or r["name"], 20), size=fs, anchor=MSO_ANCHOR.MIDDLE)
        w = max(0.03, abs(r["chgPct"]) / mx * half)
        box(s, zero if r["chgPct"] >= 0 else zero - w, y + rh * 0.25, w, rh * 0.5, fill=b_col, radius=0.3)
        text(s, 9.65, y, 1.2, rh, pct(r["chgPct"]), size=fs, bold=True, color=t_col, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 10.95, y, 1.78, rh, man(r["diff"], True), size=fs, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def slide_reasons(prs, when, movers):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "변동 사유", 2, when)
    for i, (r, query, news) in enumerate(movers[:4]):
        x, y = 0.6 + (i % 2) * 6.16, 1.35 + (i // 2) * 2.78
        t_col, _ = tone(r["chgPct"])
        box(s, x, y, 5.97, 2.6, fill=WHITE, line=LINE, radius=0.06)
        box(s, x, y + 0.18, 0.06, 0.5, fill=t_col)
        text(s, x + 0.3, y + 0.16, 3.6, 0.5, [[(cut(r.get("label") or r["name"], 16), {"bold": True, "size": 17}),
                                              (f"   비중 {r['weight']:.1f}%", {"size": 11, "color": MUTED})]])
        text(s, x + 3.9, y + 0.14, 1.85, 0.5, pct(r["chgPct"]), size=20, bold=True, color=t_col, align=PP_ALIGN.RIGHT)
        if news:
            lines = []
            for nw in news:
                lines.append([("• " + cut(nw["title"], 46), {"size": 13})])
                lines.append([("   " + (nw["source"] or ""), {"size": 10, "color": FAINT})])
            text(s, x + 0.3, y + 0.86, 5.4, 1.6, lines, spacing=3)
        else:
            text(s, x + 0.3, y + 0.86, 5.4, 1.4, f"관련 뉴스를 찾지 못했어요.\n'{query}' 흐름을 따라 움직였을 가능성이 커요.", size=12, color=MUTED)
    text(s, 0.6, 7.02, 12.13, 0.3, "많이 움직인 종목의 최근 헤드라인을 Google 뉴스에서 자동으로 모았어요. 실제 원인과 다를 수 있어요.", size=10, color=FAINT)


def slide_news(prs, when, news):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "미국 뉴스", 3, when)
    if not news:
        text(s, 0.6, 1.5, 12, 0.5, "오늘은 뉴스를 받아오지 못했어요.", size=14, color=MUTED)
        return
    rh = 5.7 / max(len(news), 1)
    for i, nw in enumerate(news):
        y = 1.3 + i * rh
        if i:
            box(s, 0.6, y - 0.06, 12.13, 0.012, fill=LINE)
        c = box(s, 0.6, y + 0.1, 0.42, 0.42, fill=CARD, radius=0.5)
        text(s, 0.6, y + 0.1, 0.42, 0.42, str(i + 1), size=13, bold=True, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        title = [(cut(nw["title"], 52), {"bold": True, "size": 15})]
        for tag in nw.get("tags", []):
            title.append((f"  #{tag}", {"size": 11, "color": DN_T}))
        meta = f"{nw['source']} · {nw['time'].strftime('%m/%d %H:%M')}" if nw.get("time") else nw["source"]
        text(s, 1.25, y + 0.06, 9.6, 0.4, [title])
        text(s, 10.9, y + 0.1, 1.83, 0.3, meta, size=10, color=FAINT, align=PP_ALIGN.RIGHT)
        lines = [[("· " + cut(l, 75), {})] for l in nw["summary"]] if len(nw["summary"]) > 1 else [[(cut(nw["summary"][0], 150), {})]]
        text(s, 1.25, y + 0.46, 11.4, rh - 0.5, lines, size=12, color=MUTED)


def viewer(when, stamp, n):
    imgs = "\n".join(f'<img src="slide-{i}.png?v={stamp}" alt="브리핑 {i}번 슬라이드" loading="lazy">' for i in range(1, n + 1))
    return f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex"><title>아침 브리핑</title>
<style>body{{margin:0;background:#F4F5F7;color:#16181D;font-family:system-ui,-apple-system,"Apple SD Gothic Neo","Malgun Gothic",sans-serif}}
main{{max-width:1100px;margin:0 auto;padding:18px 12px 40px;display:grid;gap:12px}}
h1{{font-size:18px;margin:4px 4px 0}}p{{margin:0 4px;color:#6B6E76;font-size:13px}}
img{{width:100%;height:auto;display:block;border-radius:10px;box-shadow:0 1px 3px rgba(0,0,0,.12);background:#fff}}
.dl{{display:flex;gap:8px;margin:4px}}.dl a{{flex:1;text-align:center;padding:11px;border-radius:10px;background:#16181D;color:#fff;text-decoration:none;font-size:14px}}
.dl a+a{{background:#fff;color:#16181D;border:1px solid #D5D7DC}}</style></head>
<body><main><h1>아침 브리핑</h1><p>{when}</p>
<div class="dl"><a href="briefing.pptx?v={stamp}" download>PPT 받기</a><a href="briefing.pdf?v={stamp}">PDF 보기</a></div>
{imgs}</main></body></html>
"""


def main():
    now = datetime.now(KST)
    snap, rows, total, prev_total = portfolio()
    when = f"{now.month}월 {now.day}일 ({DAYS[now.weekday()]}) {now:%H:%M} 기준"

    ranked = sorted(rows, key=lambda r: abs(r["chgPct"]), reverse=True)
    picks = [r for r in ranked if abs(r["chgPct"]) >= 1.0][:4] or ranked[:2]
    movers, used = [], set()
    for r in picks:
        q, nws = mover_news(r, used)
        movers.append((r, q, nws))
    news = us_news(now)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_changes(prs, when, rows, total, prev_total)
    slide_reasons(prs, when, movers)
    slide_news(prs, when, news)
    OUT.mkdir(exist_ok=True)
    prs.save(OUT / "briefing.pptx")

    d = total - prev_total
    meta = {
        "date": now.strftime("%Y-%m-%d"), "when": when, "stamp": now.strftime("%Y%m%d%H%M"),
        "total": round(total), "diff": round(d), "pct": round(d / prev_total * 100, 2) if prev_total else 0,
        "movers": [{"name": r.get("label") or r["name"], "chgPct": r["chgPct"]} for r, _, _ in movers],
        "headline": news[0]["title"] if news else None,
        "totalText": man(total), "diffText": man(d, True),
    }
    (OUT / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / "index.html").write_text(viewer(when, meta["stamp"], 3), encoding="utf-8")
    print(f"브리핑 생성: 총 {meta['totalText']} ({meta['diffText']}), 변동 사유 {len(movers)}개, 미국 뉴스 {len(news)}개")


if __name__ == "__main__":
    main()
